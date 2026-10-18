"""Local-reader generalisation regressions, one per defect found by evals/unseen_worlds (invented strings only).

Each test names the structural variation the reader used to miss. None of these strings comes from a supplied
mailbox or from the unseen-world fixtures verbatim; they exercise the same grammar with other words and numbers.
"""
import io
import sys
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "harness"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from factkit import Kit  # noqa: E402
from nova.extract.local import extract  # noqa: E402
from nova.extract.local_billing import (amount_before_rate, billing_line, column_field, header_field,  # noqa: E402
                                        statement_rows, table_rows)
from nova.extract.local_common import Document, iso  # noqa: E402
from nova.extract.tables import _sheet_rows  # noqa: E402
from nova.rules.common import ledger_state  # noqa: E402
from nova.world import parse_band  # noqa: E402


def pdf_payload(text, doc="batch_01/0007.eml#doc.pdf"):
    return {"doc": doc, "kind": "pdf", "pages": text if isinstance(text, list) else [text], "is_scan": False}


def workbook(rows, preface=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in list(preface) + list(rows):
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


AGREEMENT = """Air Carriage Agreement No. ZQ-PL-AGR-61
Between Zephyr Quill Air Ltd (carrier code ZQ-AU) and Pell Lumen Pty (shipper code PL-AU).
Part 1 Term
1.1 Effective date. This Agreement takes effect on 3rd March 2061.
1.2 Initial term. It expires on 2nd March 2063.
Part 2 Payment
2.1 Payment terms. Invoices are payable 45 days after the invoice date.
2.2 Late fee. No late fee applies where payment is received within 7 days after the due date. Otherwise a late fee of 1.75% per started 30-day period applies.
Exhibit A Rates
Tariff 1: shipments picked up 03.03.2061 to 31.08.2061. Tariff 2: shipments picked up 01.09.2061 to 02.03.2062.
Route Break Tariff 1 Tariff 2
PER1-SYD4 up to 39.9 kgs 7.10 7.30
PER1-SYD4 40 to 199.9 kgs 6.20 6.40
PER1-SYD4 200 kgs and above 5.55 5.70
Exhibit F Exchange rates
March 2061 91.2500
Apr 2061 91.4000
Exhibit E Billing codes
LHC Linehaul; also billed as: Freight.
Schedule 1 Authorised contacts
Tariffs Officer is not a category.
Rates, Notices Ilsa Morrow Head of Revenue ilsa.morrow@zephyrquill.example
Waivers; Payment terms: Teo Brandt, Collections Officer <teo.brandt@zephyr-
quill.example>, ZQ-AU
Signed for Zephyr Quill Air Ltd by Ilsa Morrow on 1st March 2061."""


class WeightBreakTests(unittest.TestCase):
    def test_every_inclusive_wording_and_no_guess_for_exclusive_ones(self):
        cases = {"0 to 99.9 kg": ("0", "99.9"), "250 kgs and over": ("250", None), "up to 39.9 kgs": ("0", "39.9"),
                 "100 kg +": ("100", None), "45-99.5 kg": ("45", "99.5"), "300 kg or more": ("300", None)}
        for text, (lo, hi) in cases.items():
            got = parse_band(text)
            self.assertEqual((str(got[0]), None if got[1] is None else str(got[1])), (lo, hi), text)
        self.assertEqual(parse_band("under 100 kg"), (None, None, False))
        self.assertTrue(parse_band("Minimum charge")[2])


class HeaderMeaningTests(unittest.TestCase):
    def test_ledger_and_tracking_headers_are_read_by_their_words(self):
        want = {"Doc No.": "invoice_id", "Date Paid": "paid_date", "Payment Date": "paid_date", "Paid?": "status",
                "Outstanding (INR)": "amount", "Shipment Ref": "awb", "Event Date (DD-MM-YYYY)": "date",
                "Due On": "due_date", "Bill Date": "invoice_date", "Facility": "station"}
        for header, field in want.items():
            self.assertEqual(header_field(header), field, header)

    def test_statement_columns_are_read_by_meaning(self):
        want = {"Line value": "amount", "Unit rate": "rate", "Chg wt": "qty", "Description": "desc",
                "Collected": "date", "Route": "lane", "AWB/Ref": "awb"}
        for header, field in want.items():
            self.assertEqual(column_field(header), field, header)

    def test_header_format_hint_settles_day_month_order(self):
        rows = [["Waybill No", "Event Code", "Event Date (DD-MM-YYYY)", "Hub"], ["ZQ90012345", "Picked up", "04-07-2061", "PER"]]
        blob = workbook(rows, preface=[["Milestone extract"], []])
        doc = Document({"doc": "batch_01/0002.eml#m.xlsx", "kind": "xlsx", "pages": ["x"], "is_scan": False})
        facts = table_rows(doc, blob)
        self.assertEqual(facts[0]["data"]["date"], "2061-07-04")          # 4 July, not 7 April


class LedgerWordsTests(unittest.TestCase):
    def test_status_words(self):
        self.assertEqual(ledger_state({"status": "Settled"}), "paid")
        self.assertEqual(ledger_state({"status": "Cleared"}), "paid")
        self.assertEqual(ledger_state({"status": "Unpaid"}), "open")
        self.assertEqual(ledger_state({"status": "Yes"}), "paid")
        self.assertEqual(ledger_state({"status": "No"}), "open")
        self.assertEqual(ledger_state({"status": "Open", "paid_date": "2061-05-02"}), "paid")
        self.assertEqual(ledger_state({"status": "Part paid"}), "open")
        self.assertIsNone(ledger_state({"status": "On hold"}))


class AgreementGrammarTests(unittest.TestCase):
    def setUp(self):
        self.out = extract([pdf_payload(AGREEMENT)])
        self.contract = self.out["contracts"][0]

    def test_titled_agreement_with_parts_and_ordinal_dates(self):
        self.assertEqual(self.contract["contract_id"], "ZQ-PL-AGR-61")
        self.assertEqual(self.contract["effective_date"], "2061-03-03")
        self.assertEqual(self.contract["term"]["initial_term_end"], "2063-03-02")

    def test_grace_written_as_payment_within_n_days(self):
        self.assertEqual(self.contract["payment"]["grace_days"], "7")

    def test_named_tariff_windows_price_each_column(self):
        labels = [p["label"] for p in self.contract["rate_periods"]]
        self.assertEqual(labels, ["Tariff 1", "Tariff 2"])
        self.assertEqual(self.contract["rate_periods"][1]["pickup_from"], "2061-09-01")
        rows = [(r["band"], r["period_label"], r["rate"]) for r in self.out["rate_rows"] if r["lane"] == "PER1-SYD4"]
        self.assertIn(("200 kgs and above", "Tariff 2", "5.70"), rows)
        self.assertIn(("up to 39.9 kgs", "Tariff 1", "7.10"), rows)

    def test_a_window_read_by_two_patterns_is_one_period(self):
        # Regression: the generic window pattern also matched "Period P2: pickups ...", which the older pattern had
        # already read as "P2"; three labels for two price columns dropped every rate row.
        text = ("Freight Agreement QX-RT-EXP-55\n1.1 Effective date. Effective from 2055-01-01.\n1.2 Term. Ends on "
                "2055-12-31.\nAnnex A Rates\nPeriod P1: pickups 2055-01-01 to 2055-06-30\nPeriod P2: pickups "
                "2055-07-01 to 2055-12-31\nVAL1-RIX2 0-44.5 kg 6.10 6.25")
        out = extract([pdf_payload(text)])
        self.assertEqual([p["label"] for p in out["contracts"][0]["rate_periods"]], ["P1", "P2"])
        self.assertEqual([(r["period_label"], r["rate"]) for r in out["rate_rows"]], [("P1", "6.10"), ("P2", "6.25")])

    def test_exchange_rates_by_month_name(self):
        self.assertEqual({r["month"]: r["inr_per_usd"] for r in self.out["fx_rows"]},
                         {"2061-03": "91.2500", "2061-04": "91.4000"})

    def test_codes_under_any_codes_heading(self):
        self.assertEqual([(c["code"], c["aliases"]) for c in self.out["charge_codes"]], [("LHC", ["Freight"])])

    def test_schedule_rows_in_other_layouts_and_wrapped_email(self):
        people = {p["email"]: p for p in self.out["authorised_persons"]}
        self.assertEqual(people["ilsa.morrow@zephyrquill.example"]["categories"], ["Rates", "Notices"])
        self.assertEqual(people["ilsa.morrow@zephyrquill.example"]["name"], "Ilsa Morrow")
        teo = people["teo.brandt@zephyr-quill.example"]
        self.assertEqual((teo["name"], teo["party"], teo["categories"]), ("Teo Brandt", "ZQ-AU", ["Waivers", "Payment terms"]))
        self.assertEqual(len(people), 2)


class AmendmentSentenceTests(unittest.TestCase):
    def test_rate_change_in_another_word_order(self):
        text = ("Amendment No. 3 to Air Carriage Agreement ZQ-PL-AGR-61\nThis amendment takes effect on 1st June 2061.\n"
                "1. Shipments on PER1-SYD4 in the 40 to 199.9 kgs break picked up on or after 15th June 2061 are "
                "charged USD 5.95 per kg.\nSigned for Zephyr Quill Air Ltd by Ilsa Morrow.")
        out = extract([pdf_payload(text)])
        row = next(r for r in out["rate_rows"])
        self.assertEqual((row["lane"], row["band_min_kg"], row["band_max_kg"], row["rate"]), ("PER1-SYD4", "40", "199.9", "5.95"))
        self.assertEqual(out["contracts"][0]["amendment_changes"][0]["effective_from"], "2061-06-15")


class InvoiceLayoutTests(unittest.TestCase):
    def test_amount_first_header_is_detected_from_column_order(self):
        self.assertTrue(amount_before_rate("No Value (USD) Weight Price Item Waybill Route Picked"))
        self.assertFalse(amount_before_rate("No Waybill Item Route Weight Price Value (USD)"))

    def test_amount_first_percentage_line_with_an_empty_quantity_cell(self):
        row = billing_line("4 812.44 18.5% Fuel levy 18.5% of 4,391.57 ZQ90012345", ["Fuel levy"], True)
        self.assertEqual((row["amount"], row["pct"], row["base_amount"]), ("812.44", "18.5", "4391.57"))

    def test_late_payment_charge_names_its_invoice_and_base(self):
        row = billing_line("Late payment charge - inv. ZQ-61-0042 (paid 9 days late): 1 x 1.75% on 2,400.00 "
                           "1 month 1.75% 42.00", ["Late payment charge"])
        self.assertEqual((row["refers_to_invoice"], row["base_amount"], row["amount"]), ("ZQ-61-0042", "2400.00", "42.00"))

    def test_ocr_glued_waybill_and_lane_cells(self):
        # A scan's OCR dropped the space between the waybill and lane cells (found by w4_d25).
        row = billing_line("Airfreight 176-55099001PER1-SYD4 2061-04-02 120.0 kgs 6.20 744.00", ["Airfreight"]) or {}
        self.assertEqual(tuple(row.get(k) for k in ("awb", "lane", "quantity", "rate", "amount")),
                         ("176-55099001", "PER1-SYD4", "120.0", "6.20", "744.00"))

    def test_kgs_unit_keeps_the_weight(self):
        row = billing_line("Airfreight ZQ90012345 PER1-SYD4 2061-04-02 120.0 kgs 6.20 744.00", ["Airfreight"])
        self.assertEqual((row["quantity"], row["unit"]), ("120.0", "kgs"))

    def test_statement_rows_by_header_meaning_in_any_column_order(self):
        grid = [["ZEPHYR QUILL - STATEMENT OF ACCOUNT"], ["Statement No: ZQ-ST-6104"],
                ["Line value", "Price", "Chg wt", "Route", "Collected", "Description", "AWB/Ref"],
                ["744.00", "6.20", "120.0 kg", "PER1-SYD4", "2 April 2061", "Airfreight", "ZQ90012345"],
                ["137.64", "18.5%", "", "", "", "Fuel levy 18.5% of 744.00", "ZQ90012345"]]
        doc = Document({"doc": "batch_01/0003.eml#s.xlsx", "kind": "xlsx", "is_scan": False, "grid": [grid],
                        "pages": ["\n".join(" | ".join(c for c in r if c) for r in grid)]})
        out = {"invoice_lines": []}
        self.assertTrue(statement_rows(doc, "ZQ-ST-6104", ["Airfreight", "Fuel levy"], out))
        a, b = out["invoice_lines"]
        self.assertEqual((a["amount"], a["rate"], a["quantity"], a["lane"], a["pickup_date"]),
                         ("744.00", "6.20", "120.0", "PER1-SYD4", "2061-04-02"))
        self.assertEqual((b["pct"], b["base_amount"], b["quantity"]), ("18.5", "744.00", None))


class CreditMemoTests(unittest.TestCase):
    def test_credit_memo_with_a_written_quarter(self):
        text = ("ZEPHYR QUILL AIR LTD\nCREDIT MEMO ZQ-CM-0081\nDate 12 October 2061\nAgreement ZQ-PL-AGR-61\n"
                "Reason: Quarterly volume rebate Q3 2061 (statement ZQ-ST-6109)\nCredit amount USD 880.00")
        note = extract([pdf_payload(text)])["credit_notes"][0]
        self.assertEqual((note["credit_id"], note["period"], note["credit_type"], note["amount"], note["refers_to_invoice"]),
                         ("ZQ-CM-0081", "2061-Q3", "rebate", "880.00", "ZQ-ST-6109"))


class DateTests(unittest.TestCase):
    def test_ordinal_and_dashed_numeric_dates(self):
        self.assertEqual(iso("payable by 21st May 2061"), "2061-05-21")
        self.assertEqual(iso("on 14-03-2061", "DMY"), "2061-03-14")
        self.assertIsNone(iso("on 04-03-2061"))                        # day/month order unknown: never guessed
        self.assertEqual(iso("on 04-03-2061", "DMY"), "2061-03-04")


class HoldoutFoundTests(unittest.TestCase):
    """Defects the frozen-reader holdout run exposed (w2_h05, w4_h12); fixed afterwards at the general level."""

    def test_dotted_numeric_dates_are_day_first_when_no_day_exceeds_twelve(self):
        from nova.extract.validate import PageIndex
        self.assertEqual(PageIndex._date_order(["Inv. Date 10.06.2061\nDue 01.07.2061\nCollected 05.06.2061"]), "DMY")
        self.assertIsNone(PageIndex._date_order(["Invoice date 03/04/2061"]))            # slashes: never guessed
        self.assertIsNone(PageIndex._date_order(["03/04/2061 and 05.06.2061"]))          # mixed: never guessed
        self.assertEqual(PageIndex._date_order(["04/23/2061"]), "MDY")

    def test_amendment_sentence_that_moves_the_term_end(self):
        text = ("Amendment No. 2 to Air Carriage Agreement ZQ-PL-AGR-61\nThis amendment is effective 9th September "
                "2062.\n1. Clause 1.2 is replaced: the initial term expires on 30th June 2063.\n"
                "Signed for Zephyr Quill Air Ltd by Ilsa Morrow.")
        contract = extract([pdf_payload(text)])["contracts"][0]
        self.assertEqual((contract["kind"], contract["term"]["initial_term_end"], contract["effective_date"]),
                         ("amendment", "2063-06-30", "2062-09-09"))

    def test_extension_keeps_the_agreement_renewal_terms(self):
        from nova.rules.expiry import operative_date
        k = Kit()
        try:
            k.contract(term={"initial_term_end": "2041-01-09", "auto_renews": True, "non_renewal_notice_days": "60",
                             "non_renewal_notice_deadline": "2040-11-10"})
            amd = k.doc("batch_02/0001.eml#AMD-5.pdf", "2040-10-01", "omar.lind@pellwick.example", "amendment")
            k.fact("contracts", amd, {"contract_id": "PW-OB-EXP-07/AMD-5", "kind": "amendment",
                                      "amends_contract_id": "PW-OB-EXP-07", "signed": True,
                                      "effective_date": "2040-10-01", "term": {"initial_term_end": "2041-07-31"}})
            contract = next(iter(k.world("2040-10-15").contracts.values()))
            when, _ = operative_date(contract)
            self.assertEqual(str(when), "2041-06-01")                # 60 days before the new end, still renewing
        finally:
            k.close()


class ChargeRoleTests(unittest.TestCase):
    def test_role_from_alias_when_the_name_is_the_contract_own_word(self):
        k = Kit()
        try:
            doc = k.contract()
            k.fact("charge_codes", doc, {"contract_id": "PW-OB-EXP-07", "code": "LHC", "name": "Linehaul",
                                         "aliases": ["Freight"]})
            contract = next(iter(k.world("2040-12-31").contracts.values()))
            self.assertEqual(contract.code_for("Linehaul"), "FRT")
        finally:
            k.close()


if __name__ == "__main__":
    unittest.main()
