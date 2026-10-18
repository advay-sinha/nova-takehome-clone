"""Deterministic document writers for the unseen worlds: PDF pages, XLSX sheets and .eml messages.

Byte-stable output (fixed metadata, zip entry times and message ids) so a rebuild is identical. No harness code.
"""
from __future__ import annotations

import re
import shutil
import textwrap
import zipfile
from datetime import datetime, timezone
from email.message import EmailMessage
from email.utils import format_datetime
from io import BytesIO
from pathlib import Path

import pymupdf
from openpyxl import Workbook

FIXED = datetime(2042, 6, 1)


def pdf(pages: list[list], scan: bool = False, wrap: int = 98, font: str = "helv") -> bytes:
    """pages: list of rows; a row is a str (wrapped paragraph), ("h", str) a heading, ("gap",) a blank line, or a
    list of (x, text) cells on one line."""
    doc = pymupdf.open()
    for rows in pages:
        page = doc.new_page(width=595, height=842)
        y = 54
        for row in rows:
            if isinstance(row, tuple) and row[0] == "h":
                y += 6
                page.insert_text((48, y), row[1], fontsize=11, fontname="hebo")
                y += 16
            elif isinstance(row, tuple) and row[0] == "gap":
                y += 10
            elif isinstance(row, list):
                for x, text in row:
                    page.insert_text((x, y), str(text), fontsize=8.4, fontname=font)
                y += 13
            else:
                for part in textwrap.wrap(row, wrap) or [""]:
                    page.insert_text((48, y), part, fontsize=9, fontname=font)
                    y += 12.5
                y += 3
            if y > 800:
                raise ValueError("page overflow: split the content over more pages")
    if scan:
        image_doc = pymupdf.open()
        for page in doc:
            pix = page.get_pixmap(dpi=170)
            out = image_doc.new_page(width=595, height=842)
            out.insert_image(out.rect, stream=pix.tobytes("png"))
        doc = image_doc
    doc.set_metadata({"producer": "evals/unseen_worlds", "creationDate": "D:20420601000000Z",
                      "modDate": "D:20420601000000Z"})
    return doc.tobytes(garbage=3, deflate=True, no_new_id=True)


def xlsx(sheets: list[tuple[str, list[list]]], money_cols: set[int] = frozenset(), preface: list[list] = ()) -> bytes:
    """sheets: [(title, rows)] with the header as rows[0]; `preface` rows go above the header on the first sheet."""
    wb = Workbook()
    for i, (title, rows) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        for r in (preface if i == 0 else []):
            ws.append(r)
        for r in rows:
            ws.append(r)
            for j, cell in enumerate(ws[ws.max_row]):
                if j in money_cols and isinstance(cell.value, float):
                    cell.number_format = "#,##0.00"
    wb.properties.created = wb.properties.modified = FIXED
    buf = BytesIO()
    wb.save(buf)
    out = BytesIO()
    with zipfile.ZipFile(buf) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*", rb"\g<1>2042-06-01T00:00:00Z", data)
            entry = zipfile.ZipInfo(info.filename, (2042, 6, 1, 0, 0, 0))
            entry.create_system = 0
            dst.writestr(entry, data, zipfile.ZIP_DEFLATED)
    return out.getvalue()


MIME = {"pdf": ("application", "pdf"),
        "xlsx": ("application", "vnd.openxmlformats-officedocument.spreadsheetml.sheet")}


class Mailbox:
    def __init__(self, name: str):
        self.name = name
        self.batches: dict[str, list[EmailMessage]] = {}

    def mail(self, batch: str, when: str, sender: tuple, to: tuple, subject: str, body: str, attachments=(),
             hour: int | None = None, tz_offset: str = "+0000"):
        """when: YYYY-MM-DD (receipt date in UTC unless tz_offset moves it)."""
        msg = EmailMessage()
        msg["From"] = f"{sender[0]} <{sender[1]}>"
        msg["To"] = f"{to[0]} <{to[1]}>"
        n = sum(len(v) for v in self.batches.values()) + 1
        hh = hour if hour is not None else 8 + n % 9
        stamp = datetime.fromisoformat(f"{when}T{hh:02d}:{(7 * n) % 60:02d}:00").replace(tzinfo=timezone.utc)
        msg["Date"] = format_datetime(stamp).replace("+0000", tz_offset) if tz_offset != "+0000" else format_datetime(stamp)
        msg["Subject"] = subject
        msg["Message-ID"] = f"<{self.name}-{n:03d}@{sender[1].split('@')[1]}>"
        msg.set_content(textwrap.dedent(body).strip("\n") + "\n")
        for filename, data in attachments:
            maintype, subtype = MIME[filename.rsplit(".", 1)[1]]
            msg.add_attachment(data, maintype=maintype, subtype=subtype, filename=filename)
        if msg.is_multipart():
            msg.set_boundary(f"==uw-{self.name}-{n:03d}==")
        self.batches.setdefault(batch, []).append(msg)
        return f"{batch}/{len(self.batches[batch]):04d}.eml"

    def write(self, root: Path):
        out = root / self.name
        if out.exists():
            shutil.rmtree(out)
        for batch, msgs in sorted(self.batches.items()):
            (out / batch).mkdir(parents=True)
            for i, m in enumerate(msgs, 1):
                (out / batch / f"{i:04d}.eml").write_bytes(bytes(m))
        return out
